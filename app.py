from __future__ import annotations

import hashlib
import os
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from typing import Any

import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.exc import SQLAlchemyError
from streamlit.errors import StreamlitSecretNotFoundError

from storage import load_workspace, save_workspace


APP_TITLE = "KPIling | KPI Scorecard"
STATUS_COLORS = {
    "On Track": {"text": "#18794e", "fill": "#e7f5ec"},
    "At Risk": {"text": "#9a6700", "fill": "#fff5d6"},
    "Off Track": {"text": "#b42318", "fill": "#fdecea"},
}
HEALTH_COLORS = {
    "Healthy": {"text": "#18794e", "fill": "#e7f5ec"},
    "Needs Attention": {"text": "#9a6700", "fill": "#fff5d6"},
    "Critical": {"text": "#b42318", "fill": "#fdecea"},
}

TARGET_ROWS = [
    {
        "Metric": "Revenue (MTD)",
        "Definition": "Recognized revenue recorded during the reporting period.",
        "Unit": "USD",
        "Weekly Target": 65000.0,
        "Monthly Target": 260000.0,
        "Higher is Better?": True,
        "Data Source": "Finance / billing system",
    },
    {
        "Metric": "Active Users (WAU)",
        "Definition": "Unique users active at least once during the week.",
        "Unit": "users",
        "Weekly Target": 12000.0,
        "Monthly Target": 12500.0,
        "Higher is Better?": True,
        "Data Source": "Product analytics",
    },
    {
        "Metric": "Conversion Rate",
        "Definition": "Completed conversions divided by signups.",
        "Unit": "%",
        "Weekly Target": 0.065,
        "Monthly Target": 0.07,
        "Higher is Better?": True,
        "Data Source": "Growth funnel report",
    },
    {
        "Metric": "Support Tickets Open",
        "Definition": "Unresolved support tickets at the end of the reporting week.",
        "Unit": "tickets",
        "Weekly Target": 45.0,
        "Monthly Target": 40.0,
        "Higher is Better?": False,
        "Data Source": "Support platform",
    },
    {
        "Metric": "CSAT Score",
        "Definition": "Mean customer satisfaction response on a five-point scale.",
        "Unit": "score / 5",
        "Weekly Target": 4.5,
        "Monthly Target": 4.6,
        "Higher is Better?": True,
        "Data Source": "Post-interaction survey",
    },
    {
        "Metric": "Headcount",
        "Definition": "Active employees at the end of the reporting week.",
        "Unit": "people",
        "Weekly Target": 140.0,
        "Monthly Target": 145.0,
        "Higher is Better?": True,
        "Data Source": "HRIS roster",
    },
]

METRIC_SPECS = [
    {"metric": "Revenue (MTD)", "column": "Revenue", "format": "currency", "aggregation": "sum"},
    {"metric": "Active Users (WAU)", "column": "Active Users", "format": "count", "aggregation": "mean"},
    {"metric": "Conversion Rate", "column": "Conversion Rate", "format": "percent", "aggregation": "conversion"},
    {"metric": "Support Tickets Open", "column": "Support Tickets Open", "format": "count", "aggregation": "last"},
    {"metric": "CSAT Score", "column": "CSAT", "format": "score", "aggregation": "mean"},
    {"metric": "Headcount", "column": "Headcount", "format": "count", "aggregation": "last"},
]

RAW_COLUMNS = [
    "Week Start",
    "Week End",
    "Week #",
    "Revenue",
    "Active Users",
    "Signups",
    "Conversions",
    "Conversion Rate",
    "Support Tickets Open",
    "CSAT",
    "Headcount",
    "Notes",
]


def generate_sample_data(today: date | None = None) -> pd.DataFrame:
    """Create 12 deterministic, realistic weekly snapshots ending last Sunday."""
    today = today or date.today()
    current_monday = today - timedelta(days=today.weekday())
    latest_monday = current_monday - timedelta(days=7)
    rng = np.random.default_rng(2026)
    rows: list[dict[str, Any]] = []

    for index in range(12):
        week_start = latest_monday - timedelta(weeks=11 - index)
        week_end = week_start + timedelta(days=6)
        signups = int(round(1_850 + index * 24 + rng.normal(0, 75)))
        conversions = int(round(signups * (0.058 + index * 0.0007 + rng.normal(0, 0.003))))
        revenue = round(56_000 + index * 1_180 + rng.normal(0, 3_100), 2)
        tickets = int(round(50 - index * 0.8 + rng.normal(0, 5)))
        csat = round(float(np.clip(4.18 + index * 0.025 + rng.normal(0, 0.08), 3.8, 4.9)), 2)
        rows.append(
            {
                "Week Start": pd.Timestamp(week_start),
                "Week End": pd.Timestamp(week_end),
                "Week #": week_start.isocalendar().week,
                "Revenue": revenue,
                "Active Users": int(round(9_600 + index * 205 + rng.normal(0, 280))),
                "Signups": signups,
                "Conversions": conversions,
                "Conversion Rate": conversions / signups if signups else np.nan,
                "Support Tickets Open": max(tickets, 0),
                "CSAT": csat,
                "Headcount": int(round(128 + index * 1.25 + rng.normal(0, 1.4))),
                "Notes": [
                    "Baseline week",
                    "Lifecycle email test",
                    "Referral campaign",
                    "New onboarding flow",
                    "Spring promotion",
                    "Pricing page update",
                    "Paid search refresh",
                    "Customer advisory session",
                    "Product release",
                    "Partner launch",
                    "Retention campaign",
                    "Quarterly close",
                ][index],
            }
        )

    frame = pd.DataFrame(rows, columns=RAW_COLUMNS)
    frame.loc[4, "CSAT"] = np.nan
    frame.loc[8, "Support Tickets Open"] = 96
    return frame


def prepare_uploaded_data(uploaded_data: pd.DataFrame) -> pd.DataFrame:
    """Validate weekly CSV columns and recalculate derived fields."""
    if uploaded_data.empty:
        raise ValueError("The CSV has a header but no weekly records.")

    frame = uploaded_data.copy()
    headers = [str(column).strip() for column in frame.columns]
    if len({header.casefold() for header in headers}) != len(headers):
        raise ValueError("Column names must be unique after trimming spaces.")
    frame.columns = headers
    canonical_columns = {column.casefold(): column for column in RAW_COLUMNS}
    frame = frame.rename(columns={column: canonical_columns.get(column.casefold(), column) for column in frame.columns})

    required_columns = [
        "Week Start",
        "Revenue",
        "Active Users",
        "Signups",
        "Conversions",
        "Support Tickets Open",
        "CSAT",
        "Headcount",
    ]
    missing_columns = [column for column in required_columns if column not in frame.columns]
    if missing_columns:
        raise ValueError(f"Missing required columns: {', '.join(missing_columns)}.")

    frame["Week Start"] = pd.to_datetime(frame["Week Start"], errors="coerce")
    if frame["Week Start"].isna().any():
        raise ValueError("Week Start must contain a valid date for every row (for example, 2026-09-28).")
    if "Week End" in frame.columns:
        week_end = pd.to_datetime(frame["Week End"], errors="coerce")
        frame["Week End"] = week_end.fillna(frame["Week Start"] + pd.Timedelta(days=6))
    else:
        frame["Week End"] = frame["Week Start"] + pd.Timedelta(days=6)

    numeric_columns = [
        "Revenue",
        "Active Users",
        "Signups",
        "Conversions",
        "Support Tickets Open",
        "CSAT",
        "Headcount",
    ]
    for column in numeric_columns:
        frame[column] = pd.to_numeric(frame[column], errors="coerce")

    frame["Week #"] = frame["Week Start"].dt.isocalendar().week.astype("Int64")
    frame["Conversion Rate"] = frame["Conversions"].div(frame["Signups"].where(frame["Signups"].ne(0)))
    if "Notes" not in frame.columns:
        frame["Notes"] = ""

    extra_columns = [column for column in frame.columns if column not in RAW_COLUMNS]
    return frame[RAW_COLUMNS + extra_columns].sort_values("Week Start").reset_index(drop=True)


def recalculate_derived_fields(data: pd.DataFrame) -> pd.DataFrame:
    """Keep derived weekly fields consistent after direct table edits."""
    frame = data.copy()
    frame["Week Start"] = pd.to_datetime(frame["Week Start"], errors="coerce")
    frame["Week End"] = pd.to_datetime(frame["Week End"], errors="coerce")
    frame["Week #"] = frame["Week Start"].dt.isocalendar().week.astype("Int64")
    for column in ["Revenue", "Active Users", "Signups", "Conversions", "Support Tickets Open", "CSAT", "Headcount"]:
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    frame["Conversion Rate"] = frame["Conversions"].div(frame["Signups"].where(frame["Signups"].ne(0)))
    return frame.reset_index(drop=True)


def default_targets() -> pd.DataFrame:
    return pd.DataFrame(TARGET_ROWS)


def status_for_value(value: float, target: float, higher_is_better: bool) -> str:
    """Apply the scorecard's explicit 90% / 115% guard-band rules."""
    if pd.isna(value) or pd.isna(target):
        return "Off Track"
    if higher_is_better:
        if value >= target:
            return "On Track"
        risk_floor = target * 0.9
        if value >= risk_floor or np.isclose(value, risk_floor, rtol=1e-12, atol=1e-12):
            return "At Risk"
        return "Off Track"
    if value <= target:
        return "On Track"
    risk_ceiling = target * 1.15
    if value <= risk_ceiling or np.isclose(value, risk_ceiling, rtol=1e-12, atol=1e-12):
        return "At Risk"
    return "Off Track"


def _period_frames(data: pd.DataFrame, view: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    ordered = data.sort_values("Week Start").copy()
    if view == "Weekly":
        return ordered.tail(1), ordered.iloc[-2:-1]

    latest = ordered.iloc[-1]["Week Start"]
    month_mask = ordered["Week Start"].dt.to_period("M") == latest.to_period("M")
    current = ordered.loc[month_mask]
    previous_mask = ordered["Week Start"].dt.to_period("M") == (latest.to_period("M") - 1)
    previous = ordered.loc[previous_mask]
    if current.empty:
        current = ordered.tail(4)
    if previous.empty:
        current_start = current["Week Start"].min()
        previous = ordered.loc[ordered["Week Start"] < current_start].tail(max(len(current), 1))
    return current, previous


def _aggregate_metric(frame: pd.DataFrame, spec: dict[str, str]) -> float:
    if frame.empty:
        return float("nan")
    aggregation = spec["aggregation"]
    if aggregation == "conversion":
        signups = frame["Signups"].sum(min_count=1)
        return float(frame["Conversions"].sum() / signups) if signups else float("nan")
    series = frame[spec["column"]]
    if aggregation == "sum":
        return float(series.sum(min_count=1))
    if aggregation == "mean":
        return float(series.mean())
    return float(series.iloc[-1]) if not pd.isna(series.iloc[-1]) else float("nan")


def calculate_kpis(data: pd.DataFrame, targets: pd.DataFrame, view: str) -> pd.DataFrame:
    """Calculate current value, comparison change, target, and status per KPI."""
    current_frame, previous_frame = _period_frames(data, view)
    target_column = "Weekly Target" if view == "Weekly" else "Monthly Target"
    target_lookup = targets.set_index("Metric").to_dict(orient="index")
    output: list[dict[str, Any]] = []

    for spec in METRIC_SPECS:
        target_row = target_lookup[spec["metric"]]
        current_value = _aggregate_metric(current_frame, spec)
        previous_value = _aggregate_metric(previous_frame, spec)
        target = float(target_row[target_column])
        higher_is_better = bool(target_row["Higher is Better?"])
        change = current_value - previous_value if pd.notna(previous_value) else float("nan")
        change_pct = change / abs(previous_value) if pd.notna(previous_value) and previous_value != 0 else float("nan")
        output.append(
            {
                "Metric": spec["metric"],
                "Current Value": current_value,
                "Previous Value": previous_value,
                "Change": change,
                "Change %": change_pct,
                "Target": target,
                "Higher is Better?": higher_is_better,
                "Status": status_for_value(current_value, target, higher_is_better),
                "Format": spec["format"],
            }
        )
    return pd.DataFrame(output)


def run_health_checks(data: pd.DataFrame, today: date | None = None) -> tuple[pd.DataFrame, int, str]:
    """Return check-level findings, a 0-100 score, and overall health label."""
    today = today or date.today()
    results: list[dict[str, Any]] = []

    for column in data.columns:
        missing_count = int(data[column].isna().sum())
        results.append(
            {
                "Check": f"Missing values: {column}",
                "Result": "Issue" if missing_count else "Pass",
                "Details": f"{missing_count} missing value(s) found." if missing_count else "No missing values.",
                "Affected Rows": ", ".join(str(index + 1) for index in data.index[data[column].isna()]),
                "Action": f"Review and populate {column} from its source system." if missing_count else "No action required.",
            }
        )

    duplicate_mask = data["Week Start"].duplicated(keep=False)
    duplicate_count = int(duplicate_mask.sum())
    results.append(
        {
            "Check": "Duplicate Week Start dates",
            "Result": "Issue" if duplicate_count else "Pass",
            "Details": f"{duplicate_count} row(s) have a repeated week start." if duplicate_count else "All week start dates are unique.",
            "Affected Rows": ", ".join(str(index + 1) for index in data.index[duplicate_mask]),
            "Action": "Deduplicate the weekly extract before publishing." if duplicate_count else "No action required.",
        }
    )

    ticket_mask = data["Support Tickets Open"] > 80
    ticket_count = int(ticket_mask.sum())
    results.append(
        {
            "Check": "Support tickets outlier (> 80)",
            "Result": "Issue" if ticket_count else "Pass",
            "Details": f"{ticket_count} week(s) exceed 80 open tickets." if ticket_count else "No weeks exceed 80 open tickets.",
            "Affected Rows": ", ".join(str(index + 1) for index in data.index[ticket_mask]),
            "Action": "Confirm the ticket backlog with Support; investigate the service-level breach." if ticket_count else "No action required.",
        }
    )

    csat_mask = data["CSAT"].notna() & ~data["CSAT"].between(3.5, 5.0)
    csat_count = int(csat_mask.sum())
    results.append(
        {
            "Check": "CSAT range (3.5-5.0)",
            "Result": "Issue" if csat_count else "Pass",
            "Details": f"{csat_count} non-missing score(s) are outside 3.5-5.0." if csat_count else "All reported CSAT scores are within 3.5-5.0.",
            "Affected Rows": ", ".join(str(index + 1) for index in data.index[csat_mask]),
            "Action": "Validate survey scale and source values." if csat_count else "No action required.",
        }
    )

    latest_week_end = pd.to_datetime(data["Week End"], errors="coerce").max()
    age_days = (pd.Timestamp(today) - latest_week_end.normalize()).days if pd.notna(latest_week_end) else None
    if age_days is None:
        freshness, freshness_detail = "Stale", "No valid week end date is available."
    elif age_days <= 7:
        freshness, freshness_detail = "Fresh", f"Latest week ended {age_days} day(s) ago."
    elif age_days <= 14:
        freshness, freshness_detail = "Aging", f"Latest week ended {age_days} day(s) ago."
    else:
        freshness, freshness_detail = "Stale", f"Latest week ended {age_days} day(s) ago."
    results.append(
        {
            "Check": "Data freshness",
            "Result": freshness,
            "Details": freshness_detail,
            "Affected Rows": "",
            "Action": "Refresh the weekly extract." if freshness != "Fresh" else "No action required.",
        }
    )

    checks = pd.DataFrame(results)
    issue_count = int((checks["Result"] == "Issue").sum())
    score = max(0, 100 - issue_count * 12 - (15 if freshness == "Stale" else 5 if freshness == "Aging" else 0))
    if freshness == "Stale" or duplicate_count or issue_count >= 4:
        overall = "Critical"
    elif issue_count or freshness == "Aging":
        overall = "Needs Attention"
    else:
        overall = "Healthy"
    return checks, score, overall


def _format_value(value: float, value_format: str) -> str:
    if pd.isna(value):
        return "N/A"
    if value_format == "currency":
        return f"${value:,.0f}"
    if value_format == "percent":
        return f"{value:.1%}"
    if value_format == "score":
        return f"{value:.2f} / 5.0"
    return f"{value:,.0f}"


def _change_is_favorable(change: float, higher_is_better: bool) -> bool | None:
    if pd.isna(change) or change == 0:
        return None
    return (change > 0) == higher_is_better


@st.cache_data(max_entries=4, show_spinner=False)
def build_excel_report(
    kpis: pd.DataFrame,
    raw_data: pd.DataFrame,
    health_checks: pd.DataFrame,
    health_score: int,
    health_status: str,
    targets: pd.DataFrame,
    view: str,
    refreshed_at: datetime,
) -> bytes:
    """Build a polished five-sheet Excel report using openpyxl styles."""
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        kpis.drop(columns=["Format"]).to_excel(writer, sheet_name="Dashboard KPIs", index=False, startrow=2)
        raw_data.to_excel(writer, sheet_name="Full Raw Data", index=False, startrow=2)
        health_checks.to_excel(writer, sheet_name="Data Health", index=False, startrow=2)
        targets.to_excel(writer, sheet_name="Targets & Definitions", index=False, startrow=2)

        summary = writer.book.create_sheet("Cover Summary", 0)
        summary.merge_cells("A1:F1")
        summary["A1"] = "KPILING  /  KPI SCORECARD"
        summary["A1"].font = Font(name="Aptos Display", size=22, bold=True, color="FFFFFF")
        summary["A1"].fill = PatternFill("solid", fgColor="123A36")
        summary["A1"].alignment = Alignment(vertical="center")
        summary.row_dimensions[1].height = 44
        summary["A3"] = "Report generated"
        summary["B3"] = refreshed_at.strftime("%Y-%m-%d %H:%M")
        summary["A4"] = "Reporting view"
        summary["B4"] = view
        summary["A5"] = "Overall data health"
        summary["B5"] = f"{health_status} ({health_score}/100)"
        summary["A7"] = "Scorecard snapshot"
        summary["A7"].font = Font(bold=True, color="123A36", size=13)
        summary_headers = ["Metric", "Current Value", "Target", "Status"]
        for column, header in enumerate(summary_headers, start=1):
            cell = summary.cell(row=8, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B5B")
        for row_number, record in enumerate(kpis.to_dict(orient="records"), start=9):
            summary.cell(row=row_number, column=1, value=record["Metric"])
            summary.cell(row=row_number, column=2, value=record["Current Value"])
            summary.cell(row=row_number, column=3, value=record["Target"])
            summary.cell(row=row_number, column=4, value=record["Status"])
        summary["A17"] = "Data quality findings"
        summary["A17"].font = Font(bold=True, color="123A36", size=13)
        summary["A18"] = "Check"
        summary["B18"] = "Result"
        summary["C18"] = "Details"
        for cell in summary[18][:3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B5B")
        for row_number, record in enumerate(health_checks.to_dict(orient="records"), start=19):
            summary.cell(row=row_number, column=1, value=record["Check"])
            summary.cell(row=row_number, column=2, value=record["Result"])
            summary.cell(row=row_number, column=3, value=record["Details"])
        for column, width in {"A": 34, "B": 24, "C": 62, "D": 20, "E": 18, "F": 18}.items():
            summary.column_dimensions[column].width = width
        summary.freeze_panes = "A8"

    workbook = writer.book
    sheet_titles = {
        "Dashboard KPIs": "DASHBOARD KPIs  /  " + view.upper() + " VIEW",
        "Full Raw Data": "FULL RAW DATA  /  12 WEEK HISTORY",
        "Data Health": "DATA HEALTH  /  AUTOMATED CHECKS",
        "Targets & Definitions": "TARGETS & DEFINITIONS",
    }
    for sheet_name, title in sheet_titles.items():
        sheet = workbook[sheet_name]
        last_column = max(sheet.max_column, 1)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
        title_cell = sheet.cell(row=1, column=1, value=title)
        title_cell.font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill("solid", fgColor="123A36")
        title_cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 34
        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:{get_column_letter(last_column)}{sheet.max_row}"
        for cell in sheet[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B5B")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[3].height = 30
        for row in sheet.iter_rows(min_row=4):
            for cell in row:
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F7F5")
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 1)

        for column_index in range(1, last_column + 1):
            header = str(sheet.cell(row=3, column=column_index).value or "")
            values = [len(str(sheet.cell(row=r, column=column_index).value or "")) for r in range(3, min(sheet.max_row, 80) + 1)]
            width = min(max(max(values, default=10) + 2, 12), 42)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
            if header in {"Week Start", "Week End"}:
                for row_number in range(4, sheet.max_row + 1):
                    sheet.cell(row=row_number, column=column_index).number_format = "mmm d, yyyy"
            elif header == "Revenue":
                for row_number in range(4, sheet.max_row + 1):
                    sheet.cell(row=row_number, column=column_index).number_format = '$#,##0.00;[Red]-$#,##0.00'
            elif header in {"Conversion Rate", "Change %"}:
                for row_number in range(4, sheet.max_row + 1):
                    sheet.cell(row=row_number, column=column_index).number_format = "0.0%"
            elif header in {"CSAT"}:
                for row_number in range(4, sheet.max_row + 1):
                    sheet.cell(row=row_number, column=column_index).number_format = "0.00"
            elif header in {"Current Value", "Previous Value", "Change", "Target", "Weekly Target", "Monthly Target"}:
                for row_number in range(4, sheet.max_row + 1):
                    metric = sheet.cell(row=row_number, column=1).value
                    if metric == "Revenue (MTD)":
                        number_format = '$#,##0.00;[Red]-$#,##0.00'
                    elif metric == "Conversion Rate":
                        number_format = "0.0%"
                    elif metric == "CSAT Score":
                        number_format = "0.00"
                    else:
                        number_format = "#,##0.0"
                    sheet.cell(row=row_number, column=column_index).number_format = number_format

        for row_number in range(4, sheet.max_row + 1):
            for column_index in range(1, sheet.max_column + 1):
                value = sheet.cell(row=row_number, column=column_index).value
                if value in STATUS_COLORS:
                    palette = STATUS_COLORS[value]
                    sheet.cell(row=row_number, column=column_index).fill = PatternFill("solid", fgColor=palette["fill"].lstrip("#"))
                    sheet.cell(row=row_number, column=column_index).font = Font(bold=True, color=palette["text"].lstrip("#"))
                elif value in {"Issue", "Stale", "Aging"}:
                    sheet.cell(row=row_number, column=column_index).fill = PatternFill("solid", fgColor="FDECEA" if value in {"Issue", "Stale"} else "FFF5D6")
                    sheet.cell(row=row_number, column=column_index).font = Font(bold=True, color="B42318" if value in {"Issue", "Stale"} else "9A6700")
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = "176B5B"

    summary = workbook["Cover Summary"]
    for row_number in range(9, 15):
        metric = summary.cell(row=row_number, column=1).value
        if metric == "Revenue (MTD)":
            number_format = '$#,##0.00;[Red]-$#,##0.00'
        elif metric == "Conversion Rate":
            number_format = "0.0%"
        elif metric == "CSAT Score":
            number_format = "0.00"
        else:
            number_format = "#,##0.0"
        summary.cell(row=row_number, column=2).number_format = number_format
        summary.cell(row=row_number, column=3).number_format = number_format
        status = summary.cell(row=row_number, column=4).value
        if status in STATUS_COLORS:
            palette = STATUS_COLORS[status]
            summary.cell(row=row_number, column=4).fill = PatternFill("solid", fgColor=palette["fill"].lstrip("#"))
            summary.cell(row=row_number, column=4).font = Font(bold=True, color=palette["text"].lstrip("#"))
    for row_number in range(19, summary.max_row + 1):
        result = summary.cell(row=row_number, column=2).value
        if result in {"Issue", "Stale", "Aging"}:
            summary.cell(row=row_number, column=2).fill = PatternFill("solid", fgColor="FDECEA" if result in {"Issue", "Stale"} else "FFF5D6")
            summary.cell(row=row_number, column=2).font = Font(bold=True, color="B42318" if result in {"Issue", "Stale"} else "9A6700")
    summary.sheet_view.showGridLines = False
    summary.sheet_properties.tabColor = "123A36"

    buffer.seek(0)
    buffer.truncate(0)
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _inject_styles() -> None:
    st.markdown(
        """
        <style>
        @import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@500;600;700;800&display=swap');
        :root { --ink: #18332f; --muted: #657873; --teal: #176b5b; --line: #e2ebe7; --surface: #ffffff; }
        html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
        .stApp { background: #f4f7f5; color: var(--ink); }
        [data-testid="stHeader"] { background: #f4f7f5; }
        [data-testid="stSidebar"] { background: #123a36; }
        [data-testid="stSidebar"] * { color: #edf6f2; }
        [data-testid="stSidebar"] hr { border-color: rgba(255,255,255,.18); }
        [data-testid="stSidebar"] button { background: #176b5b !important; border-color: #4b8e7e !important; }
        [data-testid="stSidebar"] button p { color: #ffffff !important; }
        [data-testid="stSidebar"] button:hover { background: #205f53 !important; }
        .stDownloadButton button { white-space: normal; min-height: 2.6rem; line-height: 1.2; }
        .block-container { padding-top: 5rem; padding-bottom: 3rem; max-width: 1440px; }
        h1, h2, h3 { font-family: 'Manrope', sans-serif; color: var(--ink); letter-spacing: 0; }
        h1 { font-size: 2rem !important; font-weight: 800 !important; }
        h2 { font-size: 1.3rem !important; }
        [data-testid="stMetric"] { background: #fff; border: 1px solid var(--line); padding: 1rem 1.1rem; border-radius: 8px; }
        div[data-testid="stTabs"] button { font-weight: 600; }
        .eyebrow { color: #568077; font-size: .72rem; font-weight: 700; text-transform: uppercase; letter-spacing: .1em; }
        .page-title { margin: .2rem 0 .15rem; font-family: 'Manrope',sans-serif; font-size: 2rem; line-height: 1.2; font-weight: 800; color: #18332f; }
        .page-subtitle { color: #657873; font-size: .95rem; margin-bottom: 1rem; }
        .kpi-card { min-height: 166px; background: #fff; border: 1px solid #e2ebe7; border-radius: 8px; padding: 16px 17px; margin: 0 0 14px; box-shadow: 0 3px 12px rgba(19,58,54,.035); }
        .kpi-label { color: #657873; font-weight: 600; font-size: .82rem; margin-bottom: 8px; }
        .kpi-value { font-family: 'Manrope',sans-serif; font-size: 1.9rem; line-height: 1.15; font-weight: 800; color: #18332f; }
        .kpi-meta { display:flex; gap:9px; align-items:center; margin-top:11px; flex-wrap:wrap; font-size:.77rem; }
        .status-pill { display:inline-flex; align-items:center; border-radius: 999px; padding: 3px 9px; font-weight:700; font-size:.72rem; }
        .target-copy { color:#657873; }
        .delta-good { color:#18794e; font-weight:700; }
        .delta-bad { color:#b42318; font-weight:700; }
        .delta-neutral { color:#657873; font-weight:700; }
        .health-banner { border-radius:8px; padding:15px 18px; display:flex; justify-content:space-between; align-items:center; gap:12px; margin:8px 0 16px; }
        .health-title { font-family:'Manrope',sans-serif; font-size:1.05rem; font-weight:800; }
        .health-detail { color:#586c67; font-size:.84rem; }
        .section-label { color:#657873; font-size:.75rem; font-weight:700; text-transform:uppercase; letter-spacing:.08em; margin: 1rem 0 .45rem; }
        @media(max-width:700px) { .block-container { padding-left:1rem; padding-right:1rem; } .kpi-card { min-height:150px; padding:13px; } .kpi-value { font-size:1.55rem; } }
        </style>
        """,
        unsafe_allow_html=True,
    )


def _render_kpi_card(record: pd.Series) -> None:
    value_text = _format_value(record["Current Value"], record["Format"])
    target_text = _format_value(record["Target"], record["Format"])
    change = record["Change"]
    change_pct = record["Change %"]
    higher_is_better = record["Higher is Better?"]
    favorable = _change_is_favorable(change, higher_is_better)
    if pd.isna(change):
        delta_text, delta_class = "No prior period", "delta-neutral"
    else:
        arrow = "▲" if change > 0 else "▼" if change < 0 else "•"
        delta_text = f"{arrow} {abs(change_pct):.1%}" if pd.notna(change_pct) else f"{arrow} {_format_value(abs(change), record['Format'])}"
        delta_class = "delta-good" if favorable else "delta-bad" if favorable is False else "delta-neutral"
    palette = STATUS_COLORS[record["Status"]]
    st.markdown(
        f"""
        <div class="kpi-card">
          <div class="kpi-label">{record['Metric']}</div>
          <div class="kpi-value">{value_text}</div>
          <div class="kpi-meta"><span class="{delta_class}">{delta_text} vs prior</span>
            <span class="status-pill" style="color:{palette['text']};background:{palette['fill']}">{record['Status']}</span></div>
          <div class="kpi-meta target-copy">Target&nbsp; {target_text}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def _make_export_button(
    kpis: pd.DataFrame,
    data: pd.DataFrame,
    checks: pd.DataFrame,
    health_score: int,
    health_status: str,
    targets: pd.DataFrame,
    view: str,
    refreshed_at: datetime,
    key: str,
    label: str,
) -> None:
    report = build_excel_report(kpis, data, checks, health_score, health_status, targets, view, refreshed_at)
    st.download_button(
        label=label,
        data=report,
        file_name=f"KPI_Scorecard_{refreshed_at:%Y-%m-%d}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
        key=key,
    )


def _oidc_is_configured(auth_settings: dict[str, Any]) -> bool:
    required_keys = ("redirect_uri", "cookie_secret", "client_id", "client_secret", "server_metadata_url")
    return bool(auth_settings) and all(auth_settings.get(key) for key in required_keys)


def _email_is_allowed(email: str, allowlist: set[str]) -> bool:
    return not allowlist or email.strip().casefold() in allowlist


def _workspace_id_for_identity(issuer: str, subject: str) -> str:
    return hashlib.sha256(f"{issuer}:{subject}".encode("utf-8")).hexdigest()


def _authentication_required(explicit_requirement: bool, database_url: str, allowlist: set[str]) -> bool:
    return explicit_requirement or bool(allowlist) or database_url.strip().casefold().startswith("postgres")


def _deployment_setting(name: str, default: str = "") -> str:
    environment_value = os.getenv(name, "").strip()
    if environment_value:
        return environment_value
    try:
        return str(st.secrets.get(name, default) or default).strip()
    except StreamlitSecretNotFoundError:
        return default


def _authenticate_user() -> tuple[str, str, bool]:
    """Use Streamlit OIDC when configured; fail closed when auth is required."""
    try:
        auth_settings = st.secrets.get("auth", {})
    except StreamlitSecretNotFoundError:
        auth_settings = {}

    auth_required = _deployment_setting("REQUIRE_AUTH").casefold() in {"1", "true", "yes"}
    allowlist = {
        email.strip().casefold()
        for email in _deployment_setting("ALLOWED_EMAILS").split(",")
        if email.strip()
    }
    oidc_configured = _oidc_is_configured(auth_settings)

    if auth_settings and not oidc_configured:
        st.error("OIDC settings are incomplete. Check the [auth] section in Streamlit secrets.")
        st.stop()
    if _authentication_required(auth_required, _deployment_setting("DATABASE_URL"), allowlist) and not oidc_configured:
        st.error("Authentication is required. Configure OIDC secrets before using this deployment.")
        st.stop()
    if not oidc_configured:
        return "demo", "Demo workspace", False

    if not st.user.is_logged_in:
        st.markdown("### Sign in to KPIling")
        st.button("Continue with your organization", on_click=st.login)
        st.stop()

    email = str(st.user.get("email", "")).strip()
    if not _email_is_allowed(email, allowlist):
        st.error("Your account is not authorized to access this scorecard.")
        st.stop()

    subject = str(st.user.get("sub", "")).strip()
    if not subject:
        st.error("The identity provider did not return a stable user ID.")
        st.stop()

    issuer = str(auth_settings["server_metadata_url"])
    owner_id = _workspace_id_for_identity(issuer, subject)
    display_name = str(st.user.get("name") or email or "Signed-in user")
    return owner_id, display_name, True


def _load_user_workspace(owner_id: str) -> None:
    """Load persisted state for the current identity or initialize its demo data."""
    try:
        workspace = load_workspace(owner_id)
    except SQLAlchemyError:
        st.error("The configured database is unavailable. Check DATABASE_URL and database connectivity.")
        st.stop()

    if workspace is None:
        raw_data = recalculate_derived_fields(generate_sample_data())
        targets = default_targets()
        refreshed_at = datetime.now().astimezone()
        source = "Generated sample"
    else:
        raw_data = recalculate_derived_fields(workspace["raw_data"])
        targets = workspace["targets"]
        refreshed_at = workspace["updated_at"]
        if refreshed_at.tzinfo is None:
            refreshed_at = refreshed_at.replace(tzinfo=timezone.utc)
        refreshed_at = refreshed_at.astimezone()
        source = "Saved workspace"

    st.session_state.workspace_owner = owner_id
    st.session_state.raw_data = raw_data
    st.session_state.targets = targets
    st.session_state.last_refreshed = refreshed_at
    st.session_state.raw_data_source = source
    st.session_state.raw_data_revision = st.session_state.get("raw_data_revision", 0) + 1
    st.session_state.upload_widget_revision = st.session_state.get("upload_widget_revision", 0) + 1
    st.session_state.loaded_csv_digest = None


def _save_user_workspace(owner_id: str) -> None:
    try:
        save_workspace(owner_id, st.session_state.raw_data, st.session_state.targets)
    except SQLAlchemyError:
        st.error("The change could not be saved. Check DATABASE_URL and database connectivity.")
        st.stop()


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, page_icon="N", layout="wide", initial_sidebar_state="expanded")
    _inject_styles()

    owner_id, user_label, authenticated = _authenticate_user()
    if st.session_state.get("workspace_owner") != owner_id:
        _load_user_workspace(owner_id)

    with st.sidebar:
        st.markdown("<div style='font-family:Manrope,sans-serif;font-size:1.25rem;font-weight:800'>KPILING</div>", unsafe_allow_html=True)
        st.caption("KPI Scorecard & Data Health")
        st.markdown("A compact weekly operating view for performance, data quality, and reporting.")
        if authenticated:
            st.caption(f"Signed in as {user_label}")
            if st.button("Log out", width="stretch"):
                st.logout()
        st.divider()
        st.markdown(f"**Last refreshed**  \n{st.session_state.last_refreshed:%b %d, %Y · %H:%M %Z}")
        if st.button("↻ Restore sample data", width="stretch"):
            st.session_state.raw_data = recalculate_derived_fields(generate_sample_data())
            st.session_state.raw_data_revision += 1
            st.session_state.upload_widget_revision += 1
            st.session_state.loaded_csv_digest = None
            st.session_state.raw_data_source = "Generated sample"
            st.session_state.last_refreshed = datetime.now().astimezone()
            _save_user_workspace(owner_id)
            st.rerun()
        st.markdown("<div style='height:8px'></div>", unsafe_allow_html=True)

    data = st.session_state.raw_data.copy()
    targets = st.session_state.targets.copy()
    health_checks, health_score, health_status = run_health_checks(data)
    refreshed_at = st.session_state.last_refreshed

    st.markdown('<div class="eyebrow">OPERATING REVIEW  /  WEEKLY PERFORMANCE</div>', unsafe_allow_html=True)
    st.markdown('<div class="page-title">KPI Scorecard</div>', unsafe_allow_html=True)
    period_end = pd.to_datetime(data["Week End"]).max().strftime("%b %d, %Y")
    st.markdown(f'<div class="page-subtitle">12-week operating snapshot · latest complete week ending {period_end}</div>', unsafe_allow_html=True)

    view = st.radio("Reporting period", ["Weekly", "Monthly"], horizontal=True, key="reporting_view")
    kpis = calculate_kpis(data, targets, view)
    with st.sidebar:
        _make_export_button(kpis, data, health_checks, health_score, health_status, targets, view, refreshed_at, "sidebar_export", "📥 Export Full Report")
        st.caption("Excel report · 5 formatted sheets")

    tabs = st.tabs(["Dashboard", "Raw Data", "Data Health", "Targets & Definitions", "Documentation"])

    with tabs[0]:
        heading_col, export_col = st.columns([1, 1])
        with heading_col:
            st.markdown("### Performance at a glance")
            st.caption("Period-to-period movement, current target, and delivery status.")
        with export_col:
            _make_export_button(kpis, data, health_checks, health_score, health_status, targets, view, refreshed_at, "dashboard_export", "📥 Export Full Report to Excel")
        for start in (0, 3):
            cols = st.columns(3, gap="medium")
            for col, (_, record) in zip(cols, kpis.iloc[start : start + 3].iterrows()):
                with col:
                    _render_kpi_card(record)

        status_summary = kpis[["Metric", "Current Value", "Target", "Change %", "Status"]].copy()
        status_summary["Current Value"] = [
            _format_value(row["Current Value"], row["Format"]) for _, row in kpis.iterrows()
        ]
        status_summary["Target"] = [_format_value(row["Target"], row["Format"]) for _, row in kpis.iterrows()]
        status_summary["Change vs prior"] = status_summary.pop("Change %").map(lambda value: f"{value:+.1%}" if pd.notna(value) else "N/A")
        st.markdown('<div class="section-label">Status summary</div>', unsafe_allow_html=True)
        st.dataframe(
            status_summary[["Metric", "Current Value", "Target", "Change vs prior", "Status"]],
            width="stretch",
            hide_index=True,
            column_config={"Status": st.column_config.TextColumn("Status", help="Based on the target and directionality rule.")},
        )

        st.markdown('<div class="section-label">12-week trend</div>', unsafe_allow_html=True)
        trend_metric = st.selectbox("Trend metric", ["Revenue", "Active Users", "Conversion Rate", "Support Tickets Open", "CSAT", "Headcount"], label_visibility="collapsed")
        trend = data.sort_values("Week Start")
        fig = px.line(trend, x="Week Start", y=trend_metric, markers=True, template="plotly_white")
        fig.update_traces(line_color="#176b5b", line_width=2.5, marker=dict(size=7, color="#ed9b40"), hovertemplate="%{x|%b %d}<br>%{y:,.2f}<extra></extra>")
        fig.update_layout(margin=dict(l=8, r=8, t=10, b=8), height=290, paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", font=dict(family="DM Sans, sans-serif", color="#526a64"), xaxis_title=None, yaxis_title=None)
        fig.update_xaxes(showgrid=False)
        fig.update_yaxes(gridcolor="#e6eeea", zeroline=False)
        st.plotly_chart(fig, width="stretch", config={"displayModeBar": False})

    with tabs[1]:
        st.markdown("### Weekly source data")
        st.caption("Upload a weekly CSV or edit cells directly. Week # and Conversion Rate are calculated automatically.")
        uploaded_file = st.file_uploader(
            "Upload weekly data",
            type=["csv"],
            key=f"weekly_csv_upload_{st.session_state.upload_widget_revision}",
            help="Required: Week Start, Revenue, Active Users, Signups, Conversions, Support Tickets Open, CSAT, Headcount.",
        )

        if uploaded_file is not None:
            csv_bytes = uploaded_file.getvalue()
            csv_digest = hashlib.sha256(csv_bytes).hexdigest()
            if csv_digest != st.session_state.loaded_csv_digest:
                try:
                    uploaded_data = prepare_uploaded_data(pd.read_csv(BytesIO(csv_bytes)))
                except (pd.errors.EmptyDataError, pd.errors.ParserError, UnicodeDecodeError, ValueError) as error:
                    st.error(f"CSV not loaded: {error}")
                else:
                    st.session_state.raw_data = uploaded_data
                    st.session_state.raw_data_revision += 1
                    st.session_state.loaded_csv_digest = csv_digest
                    st.session_state.raw_data_source = uploaded_file.name
                    st.session_state.last_refreshed = datetime.now().astimezone()
                    _save_user_workspace(owner_id)
                    st.rerun()
            else:
                st.success(f"Loaded {len(st.session_state.raw_data)} weekly records from {uploaded_file.name}.")

        editable_data = st.data_editor(
            st.session_state.raw_data,
            width="stretch",
            hide_index=True,
            num_rows="dynamic",
            height=460,
            key=f"raw_data_editor_{st.session_state.raw_data_revision}",
            column_config={
                "Week Start": st.column_config.DateColumn(format="MMM D, YYYY"),
                "Week End": st.column_config.DateColumn(format="MMM D, YYYY"),
                "Week #": st.column_config.NumberColumn(format="%d", disabled=True),
                "Revenue": st.column_config.NumberColumn(format="$%,.2f"),
                "Active Users": st.column_config.NumberColumn(format="%,d"),
                "Signups": st.column_config.NumberColumn(format="%,d"),
                "Conversions": st.column_config.NumberColumn(format="%,d"),
                "Support Tickets Open": st.column_config.NumberColumn(format="%,d"),
                "CSAT": st.column_config.NumberColumn(format="%.2f"),
                "Headcount": st.column_config.NumberColumn(format="%,d"),
                "Notes": st.column_config.TextColumn(width="large"),
                "Conversion Rate": st.column_config.NumberColumn(format="percent", disabled=True),
            },
        )
        if editable_data.empty:
            st.warning("Keep at least one weekly row in the scorecard.")
            st.session_state.raw_data_revision += 1
            st.rerun()
        updated_data = recalculate_derived_fields(editable_data)
        if not updated_data.equals(st.session_state.raw_data):
            st.session_state.raw_data = updated_data
            st.session_state.raw_data_revision += 1
            st.session_state.raw_data_source = "Edited in app"
            st.session_state.last_refreshed = datetime.now().astimezone()
            _save_user_workspace(owner_id)
            st.rerun()
        st.caption(f"{len(st.session_state.raw_data)} weekly rows · Source: {st.session_state.raw_data_source} · Changes update the scorecard and health checks.")

    with tabs[2]:
        st.markdown("### Data health monitor")
        health_palette = HEALTH_COLORS[health_status]
        st.markdown(
            f"""<div class="health-banner" style="background:{health_palette['fill']};border:1px solid {health_palette['fill']}">
              <div><div class="health-title" style="color:{health_palette['text']}">{health_status}</div>
              <div class="health-detail">Overall health score · {health_score}/100</div></div>
              <div class="status-pill" style="color:{health_palette['text']};background:#ffffff">{health_score} / 100</div>
            </div>""",
            unsafe_allow_html=True,
        )
        missing = data.isna().sum().rename("Missing Values")
        missing.index.name = "Column"
        missing = missing.reset_index()
        left, right = st.columns([1, 2])
        with left:
            st.markdown("**Missing values by column**")
            if missing.empty:
                st.success("No missing values found.")
            else:
                st.dataframe(missing, hide_index=True, width="stretch")
        with right:
            st.markdown("**Automated checks**")
            st.dataframe(
                health_checks[["Check", "Result", "Details", "Affected Rows"]],
                width="stretch",
                hide_index=True,
                column_config={"Result": st.column_config.TextColumn("Result", width="small")},
            )
        st.markdown("**Action log**")
        actions = health_checks.loc[health_checks["Result"].isin(["Issue", "Aging", "Stale"]), ["Check", "Details", "Action"]]
        if actions.empty:
            actions = pd.DataFrame([{"Check": "All clear", "Details": "No follow-up is needed.", "Action": "Continue routine monitoring."}])
        st.dataframe(actions, width="stretch", hide_index=True)
        st.caption("Fresh ≤ 7 days · Aging 8-14 days · Stale > 14 days. Missing values, duplicate dates, and outliers count against the score.")

    with tabs[3]:
        st.markdown("### Targets & metric definitions")
        st.caption("Edit targets and directionality below. Changes are saved to your user workspace and applied immediately.")
        edited_targets = st.data_editor(
            st.session_state.targets,
            width="stretch",
            hide_index=True,
            num_rows="fixed",
            key="targets_editor",
            column_config={
                "Metric": st.column_config.TextColumn(disabled=True),
                "Definition": st.column_config.TextColumn(width="large", disabled=True),
                "Weekly Target": st.column_config.NumberColumn(format="%.3f", min_value=0.0),
                "Monthly Target": st.column_config.NumberColumn(format="%.3f", min_value=0.0),
                "Higher is Better?": st.column_config.CheckboxColumn(),
                "Data Source": st.column_config.TextColumn(disabled=True),
            },
        )
        if not edited_targets.equals(st.session_state.targets):
            st.session_state.targets = edited_targets
            _save_user_workspace(owner_id)
            st.rerun()
        st.info("Higher-is-better: On Track at or above target, At Risk at or above 90% of target, otherwise Off Track. Lower-is-better: On Track at or below target, At Risk through 115% of target, otherwise Off Track.")

    with tabs[4]:
        st.markdown("### How to use this scorecard")
        st.markdown(
            """
            **Dashboard**

            Choose a weekly snapshot or a monthly roll-up. Weekly values compare with the previous complete week; monthly values compare with the prior calendar month. Revenue sums period totals, conversion rate is recomputed from total conversions and signups, CSAT and WAU are averaged, and open tickets and headcount use the latest snapshot.

            **Raw Data**

            Upload a CSV with the required weekly columns, or edit values directly in the table. Add or remove rows with the table controls. Week # and Conversion Rate are derived automatically, and changes update the scorecard and health checks immediately. Use **Restore sample data** in the sidebar to return to the demo dataset.

            **Data Health**

            The checks identify missing values, duplicate week starts, ticket counts above 80, CSAT outside 3.5-5.0, and stale reporting periods. Each finding includes affected rows and a suggested action.

            **Targets & Definitions**

            Adjust numeric targets or the directionality switch in the editable table. Changes apply immediately and are saved to your user workspace.

            **Deployment and data access**

            Local development uses a SQLite file under `data/`. Hosted deployments should set `DATABASE_URL` to a managed PostgreSQL database, `REQUIRE_AUTH=true`, and configure OIDC in Streamlit secrets. `ALLOWED_EMAILS` can restrict access to a comma-separated list of email addresses. Each signed-in user receives an isolated workspace.

            **CSV format**

            Required columns: `Week Start`, `Revenue`, `Active Users`, `Signups`, `Conversions`, `Support Tickets Open`, `CSAT`, and `Headcount`. Use one row per week and a date such as `2026-09-28` for Week Start. `Week End`, `Week #`, `Conversion Rate`, and `Notes` are optional; the app derives the first three and leaves missing Notes blank. Numeric blanks stay blank so Data Health can flag them. The sample intentionally contains one missing CSAT value and one 96-ticket outlier to demonstrate the checks.

            **Excel report**

            Use either export button to download a timestamped workbook with a summary, KPI results, raw data, data-health checks, and target definitions. The workbook is formatted with filters, frozen headers, readable widths, number formats, and status colors.
            """
        )
        st.markdown("**Refresh timestamp**")
        st.write(refreshed_at.strftime("%A, %B %d, %Y at %H:%M %Z"))


if __name__ == "__main__":
    main()
